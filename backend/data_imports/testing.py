"""Builders for import-batch and legacy-record tests.

Every workbook is synthetic and built in memory - the real historical
workbook must never enter this repository. Uploaders and source locators are
numbered, so any test can build as many batches and records as it needs
without colliding on a unique email or locator.
"""

from __future__ import annotations

import itertools
from collections.abc import Callable
from io import BytesIO

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl.workbook.workbook import Workbook

from accounts.models import User
from data_imports.models import ImportBatch
from data_imports.services import create_import_batch
from gradings.legacy_services import (
    link_legacy_record_to_student,
)
from gradings.models import LegacyGradingRecord, LegacyRecordKind
from students.models import Student

MAIN_TABLE_HEADER = ["Name", "Current Rank", "Testing For", "Outcome", "Comments"]

_sequence = itertools.count(1)


def _administrator() -> User:
    return User.objects.create_user(
        email=f"legacy-fixture-admin-{next(_sequence)}@example.com",
        role=User.Role.ADMINISTRATOR,
    )


def workbook_bytes(build: Callable[[Workbook], None]) -> bytes:
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    build(workbook)
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def standard_workbook(main_rows: list[list[object]], sheet: str = "06-2025") -> bytes:
    """An "Information" cover sheet plus one cycle sheet, as the real one has."""

    def build(workbook: Workbook) -> None:
        info = workbook.create_sheet("Information")
        info.append(["This workbook contains historical grading records."])
        ws = workbook.create_sheet(sheet)
        ws.append(MAIN_TABLE_HEADER)
        for row in main_rows:
            ws.append(row)

    return workbook_bytes(build)


def legacy_batch(
    content: bytes | None = None,
    uploaded_by: User | None = None,
) -> ImportBatch:
    payload = content or standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", ""]],
    )
    result = create_import_batch(
        kind=ImportBatch.Kind.LEGACY_GRADINGS,
        upload=SimpleUploadedFile("workbook.xlsx", payload),
        uploaded_by=uploaded_by or _administrator(),
    )
    return result.batch


def legacy_record(
    *,
    student: Student | None = None,
    **overrides: object,
) -> LegacyGradingRecord:
    """One imported record; linked to `student` when one is given."""
    number = next(_sequence)
    data: dict[str, object] = {
        "record_kind": LegacyRecordKind.GRADING_RESULT,
        "source_sheet": "06-2025",
        "source_locator": f"06-2025!R{number}",
        "cycle_year": 2025,
        "cycle_month": 6,
        "source_name": str(student) if student else "A. Example",
    }
    data.update(overrides)
    if "source_batch" not in data:
        data["source_batch"] = legacy_batch()
    # Saved one at a time, through the model's own full validation.
    record = LegacyGradingRecord.objects.create(**data)
    if student is None:
        return record
    return link_legacy_record_to_student(record, student, actor=_administrator())
