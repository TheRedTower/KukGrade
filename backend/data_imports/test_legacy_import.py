"""Historical workbook preview/commit: rank suggestions, adapter warnings, and
the LegacyGradingRecord commit path, on top of the adapter's own tests."""

from __future__ import annotations

import pytest
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl.workbook import Workbook
from rest_framework.test import APIClient

from audit.models import AuditAction, AuditEvent
from data_imports.models import ImportBatch
from data_imports.services import (
    cancel_import_batch,
    commit_legacy_grading_import,
    preview_legacy_grading_import,
)
from data_imports.testing import (
    MAIN_TABLE_HEADER,
    legacy_batch,
    standard_workbook,
    workbook_bytes,
)
from data_imports.tests import admin_user, instructor_user, rank
from gradings.models import LegacyGradingRecord, LegacyOutcome, LegacyRecordKind

pytestmark = pytest.mark.django_db


# --- preview -----------------------------------------------------------


def test_preview_reports_the_adapters_structural_warnings() -> None:
    # A sheet the adapter cannot place is skipped, not silently lost: the
    # administrator must see it before committing.
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(MAIN_TABLE_HEADER)
        ws.append(["A. Example", "White", "Yellow Stripe", "Pass", ""])
        workbook.create_sheet("Scratch notes").append(["Not a grading cycle"])

    preview = preview_legacy_grading_import(legacy_batch(workbook_bytes(build)))

    assert preview.total == 1
    assert preview.warnings == ["Unknown sheet 'Scratch notes' was not imported."]


def test_an_over_long_cell_is_refused_at_upload_naming_where_it_is() -> None:
    # Source text is kept verbatim, so it cannot be cut to fit; the commit
    # used to fail the whole batch naming no row, after a clean preview.
    from django.core.exceptions import ValidationError

    long_name = "N" * 201
    with pytest.raises(ValidationError) as refusal:
        legacy_batch(
            standard_workbook([[long_name, "White", "Yellow Stripe", "Pass", ""]])
        )

    message = refusal.value.message_dict["file"][0]
    assert message.startswith("06-2025!R2: source name is 201 characters")
    assert "at most 200" in message
    assert not ImportBatch.objects.exists()


def test_preview_counts_rows_by_record_kind() -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(MAIN_TABLE_HEADER)
        ws.append(["A. Example", "White", "Yellow Stripe", "Pass", ""])
        ws.append(["B. Second", "White", "Yellow Stripe", "DNA", ""])
        ws.append(["Missing"])
        ws.append(["C. Third", "note"])

    batch = legacy_batch(workbook_bytes(build))
    preview = preview_legacy_grading_import(batch)

    assert preview.total == 3
    assert preview.grading_result == 1
    assert preview.did_not_attend == 1
    assert preview.auxiliary_note == 1


def test_preview_suggests_an_exact_rank_match() -> None:
    # The seeded catalogue's Rank.name for "YELLOW" is exactly "Yellow" (see
    # ranks/migrations/0002_seed_rank_catalogue.py) - an exact match, not a
    # fuzzy one.
    content = standard_workbook(
        [["A. Example", "Yellow", "Blue Stripe", "Pass", ""]],
    )
    batch = legacy_batch(content)

    preview = preview_legacy_grading_import(batch)

    suggestion = next(s for s in preview.rank_suggestions if s.source_text == "Yellow")
    assert suggestion.suggested_rank_code == "YELLOW"


def test_preview_does_not_suggest_a_near_miss_no_fuzzy_matching() -> None:
    # "Yellow Belt" is a very plausible near-miss for the "Yellow" rank, but
    # exact-match-only means it is deliberately left unsuggested rather than
    # guessed (import Decision 4 - only the admin resolves ambiguous rank text).
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", ""]],
    )
    batch = legacy_batch(content)

    preview = preview_legacy_grading_import(batch)

    suggestion = next(
        s for s in preview.rank_suggestions if s.source_text == "Yellow Belt"
    )
    assert suggestion.suggested_rank_code is None


def test_preview_leaves_an_unmatched_rank_text_unsuggested() -> None:
    content = standard_workbook([["A. Example", "DNB", "Blue Stripe?", "Pass", ""]])
    batch = legacy_batch(content)

    preview = preview_legacy_grading_import(batch)

    suggestion = next(s for s in preview.rank_suggestions if s.source_text == "DNB")
    assert suggestion.suggested_rank_code is None


def test_preview_counts_unresolved_ranks_until_a_mapping_is_supplied() -> None:
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", ""]],
    )
    batch = legacy_batch(content)

    unmapped = preview_legacy_grading_import(batch)
    assert unmapped.unresolved_current_rank == 1

    yellow = rank("YELLOW")
    mapped = preview_legacy_grading_import(
        batch,
        rank_mappings={"Yellow Belt": yellow.code},
    )
    assert mapped.unresolved_current_rank == 0


def test_preview_counts_unresolved_free_form_outcomes() -> None:
    content = standard_workbook(
        [["A. Example", "White", "Yellow Stripe", "Testing? JKN Adi: No", ""]],
    )
    batch = legacy_batch(content)

    preview = preview_legacy_grading_import(batch)

    assert preview.unresolved_outcome == 1


def test_preview_is_recomputed_fresh_and_deterministic() -> None:
    batch = legacy_batch()
    first = preview_legacy_grading_import(batch)
    second = preview_legacy_grading_import(batch)
    assert first == second


# --- commit --------------------------------------------------------------


def test_commit_creates_one_legacy_record_per_row() -> None:
    admin = admin_user()
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", "Nice form"]],
    )
    batch = legacy_batch(content, uploaded_by=admin)
    events_before = AuditEvent.objects.count()

    committed = commit_legacy_grading_import(batch, actor=admin)

    assert committed.status == ImportBatch.Status.COMMITTED
    record = LegacyGradingRecord.objects.get(source_batch=committed)
    assert record.source_name == "A. Example"
    assert record.outcome == LegacyOutcome.PASS
    assert record.comments == "Nice form"
    assert record.current_rank is None  # no mapping supplied
    import_event = AuditEvent.objects.get(
        action=AuditAction.IMPORT_BATCH_COMMITTED,
        entity_type="ImportBatch",
        entity_id=str(batch.pk),
    )
    assert import_event.metadata["record_count"] == 1
    assert "A. Example" not in import_event.summary
    assert AuditEvent.objects.count() == events_before + 1


def test_commit_applies_approved_rank_mappings() -> None:
    admin = admin_user()
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", ""]],
    )
    batch = legacy_batch(content, uploaded_by=admin)
    yellow = rank("YELLOW")

    commit_legacy_grading_import(
        batch,
        actor=admin,
        rank_mappings={"Yellow Belt": yellow.code},
    )

    record = LegacyGradingRecord.objects.get(source_batch=batch)
    assert record.current_rank_id == yellow.pk
    assert record.source_current_rank_text == "Yellow Belt"
    # Left unresolved since no mapping was supplied for it.
    assert record.testing_for_rank is None
    assert record.source_testing_for_rank_text == "Blue Stripe"


def test_commit_never_creates_a_native_promotion_or_outcome() -> None:
    from gradings.models import CandidateOutcome, Promotion

    admin = admin_user()
    content = standard_workbook([["A. Example", "White", "Yellow Stripe", "Pass", ""]])
    batch = legacy_batch(content, uploaded_by=admin)
    promotions_before = Promotion.objects.count()
    outcomes_before = CandidateOutcome.objects.count()

    commit_legacy_grading_import(batch, actor=admin)

    assert Promotion.objects.count() == promotions_before
    assert CandidateOutcome.objects.count() == outcomes_before


def test_commit_did_not_attend_row_has_no_outcome() -> None:
    admin = admin_user()
    content = standard_workbook([["A. Example", "White", "Yellow Stripe", "DNA", ""]])
    batch = legacy_batch(content, uploaded_by=admin)

    commit_legacy_grading_import(batch, actor=admin)

    record = LegacyGradingRecord.objects.get(source_batch=batch)
    assert record.record_kind == LegacyRecordKind.DID_NOT_ATTEND
    assert record.outcome is None


def test_commit_is_not_idempotent_a_second_attempt_is_refused() -> None:
    admin = admin_user()
    batch = legacy_batch(uploaded_by=admin)
    commit_legacy_grading_import(batch, actor=admin)

    with pytest.raises(Exception, match="cannot be committed again"):
        commit_legacy_grading_import(batch, actor=admin)


def test_commit_deletes_the_temporary_source_file() -> None:
    admin = admin_user()
    batch = legacy_batch(uploaded_by=admin)
    assert batch.source_file

    committed = commit_legacy_grading_import(batch, actor=admin)

    assert not committed.source_file


def test_cancel_a_legacy_batch_deletes_temp_file() -> None:
    admin = admin_user()
    batch = legacy_batch(uploaded_by=admin)

    cancelled = cancel_import_batch(batch, actor=admin)

    assert cancelled.status == ImportBatch.Status.CANCELLED
    assert not cancelled.source_file


# --- API/permissions -----------------------------------------------------


def test_instructor_cannot_preview_or_commit_a_legacy_batch() -> None:
    client = APIClient()
    client.force_authenticate(user=instructor_user())
    batch = legacy_batch()

    preview_response = client.post(
        f"/api/data-imports/batches/{batch.pk}/preview/",
        {"rank_mappings": {}},
        format="json",
    )
    assert preview_response.status_code == 403

    commit_response = client.post(
        f"/api/data-imports/batches/{batch.pk}/commit/",
        {"rank_mappings": {}},
        format="json",
    )
    assert commit_response.status_code == 403


def test_admin_completes_the_full_legacy_flow_via_the_api() -> None:
    client = APIClient()
    admin = admin_user()
    client.force_authenticate(user=admin)
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", ""]],
    )

    create_response = client.post(
        "/api/data-imports/batches/",
        {
            "kind": ImportBatch.Kind.LEGACY_GRADINGS,
            "file": SimpleUploadedFile("workbook.xlsx", content),
        },
        format="multipart",
    )
    assert create_response.status_code == 201
    batch_id = create_response.json()["id"]

    preview_response = client.post(
        f"/api/data-imports/batches/{batch_id}/preview/",
        {"rank_mappings": {}},
        format="json",
    )
    assert preview_response.status_code == 200
    body = preview_response.json()
    assert body["total"] == 1
    assert body["grading_result"] == 1
    assert any(s["source_text"] == "Yellow Belt" for s in body["rank_suggestions"])

    yellow = rank("YELLOW")
    commit_response = client.post(
        f"/api/data-imports/batches/{batch_id}/commit/",
        {"rank_mappings": {"Yellow Belt": yellow.code}},
        format="json",
    )
    assert commit_response.status_code == 200
    assert commit_response.json()["id"] == batch_id
    assert ImportBatch.objects.get(pk=batch_id).status == ImportBatch.Status.COMMITTED

    record = LegacyGradingRecord.objects.get(source_batch_id=batch_id)
    assert record.current_rank_id == yellow.pk
