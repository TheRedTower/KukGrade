"""Every upload is bounded while it is read, not after (KG-A10).

Limits are patched down where the real ones would need a large fixture: the
boundary logic is the same at any size. Synthetic workbooks only.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Callable
from unittest.mock import patch

import pytest
from django.conf import settings
from django.core.exceptions import ValidationError
from django.core.files.uploadedfile import SimpleUploadedFile
from openpyxl.workbook import Workbook

from data_imports.ksw_adapter import (
    WorkbookStructureError,
    parse_ksw_colchester_workbook,
)
from data_imports.models import ImportBatch
from data_imports.parsers import (
    MAX_COLUMNS,
    UnsupportedImportFileError,
    load_tabular_dataset,
)
from data_imports.services import create_import_batch
from data_imports.storage import import_upload_storage
from data_imports.testing import MAIN_TABLE_HEADER, standard_workbook, workbook_bytes
from data_imports.tests import CLEAN_ROW, admin_user, student_workbook

pytestmark = pytest.mark.django_db


def stored_files() -> set[str]:
    try:
        return set(import_upload_storage().listdir("")[1])
    except FileNotFoundError:
        return set()


def refuse_upload(kind: str, filename: str, content: bytes, match: str) -> None:
    """The upload is refused, and leaves neither a batch nor a file behind."""
    before = stored_files()
    with pytest.raises(ValidationError, match=match):
        create_import_batch(
            kind=kind,
            upload=SimpleUploadedFile(filename, content),
            uploaded_by=admin_user(),
        )
    assert not ImportBatch.objects.exists()
    assert stored_files() == before


def rows_of(count: int) -> list[list[object]]:
    return [
        [f"{index:06d}", "Alex", "Kim", "2012-05-01", "White", "ACTIVE"]
        for index in range(count)
    ]


# --- upload size ---------------------------------------------------------


class CountingFile(io.BytesIO):
    def __init__(self, content: bytes) -> None:
        super().__init__(content)
        self.bytes_read: list[int] = []

    def read(self, size: int | None = -1) -> bytes:
        data = super().read(size)
        self.bytes_read.append(len(data))
        return data


def test_the_read_is_capped_even_when_the_declared_size_is_wrong() -> None:
    limit = settings.IMPORT_UPLOAD_MAX_BYTES
    source = CountingFile(b"a" * (limit + 1024 * 1024))
    upload = SimpleUploadedFile("roster.csv", b"")
    upload.file = source
    upload.size = 10

    with pytest.raises(ValidationError, match="MiB limit"):
        create_import_batch(
            kind=ImportBatch.Kind.STUDENT_ROSTER,
            upload=upload,
            uploaded_by=admin_user(),
        )

    assert source.bytes_read == [limit + 1]
    assert not ImportBatch.objects.exists()


# --- roster rows and columns ---------------------------------------------


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_rows_up_to_the_limit_are_read_and_one_more_is_refused(
    file_format: str,
) -> None:
    filename = f"roster.{file_format}"
    with patch("data_imports.parsers.MAX_ROWS", 3):
        at_limit = student_workbook(rows_of(3), file_format=file_format)
        assert len(load_tabular_dataset(content=at_limit, filename=filename).rows) == 3

        over = student_workbook(rows_of(4), file_format=file_format)
        with pytest.raises(UnsupportedImportFileError, match="Too many rows"):
            load_tabular_dataset(content=over, filename=filename)


@pytest.mark.parametrize("file_format", ["csv", "xlsx"])
def test_columns_up_to_the_limit_are_read_and_one_more_is_refused(
    file_format: str,
) -> None:
    filename = f"roster.{file_format}"
    headers = [f"Column {index}" for index in range(MAX_COLUMNS)]
    at_limit = student_workbook([], headers=headers, file_format=file_format)
    assert len(load_tabular_dataset(content=at_limit, filename=filename).headers) == 50

    over = student_workbook(
        [],
        headers=[*headers, "One too many"],
        file_format=file_format,
    )
    with pytest.raises(UnsupportedImportFileError, match="Too many columns"):
        load_tabular_dataset(content=over, filename=filename)


def test_a_sparse_sheet_is_refused_without_reading_it_all() -> None:
    # A few kilobytes that read as 200,000 rows: the budget stops the read.
    def build(workbook: Workbook) -> None:
        sheet = workbook.create_sheet("Roster")
        sheet.append(["WKSA ID", "First Name"])
        sheet["A200000"] = "far away"

    content = workbook_bytes(build)
    assert len(content) < 10_000
    refuse_upload(
        ImportBatch.Kind.STUDENT_ROSTER,
        "roster.xlsx",
        content,
        "Too many rows",
    )


def test_cells_far_past_the_last_column_are_never_read() -> None:
    def build(workbook: Workbook) -> None:
        sheet = workbook.create_sheet("Roster")
        sheet.append(["WKSA ID", "First Name"])
        sheet["XFD2"] = "far right"

    dataset = load_tabular_dataset(
        content=workbook_bytes(build),
        filename="roster.xlsx",
    )

    assert dataset.headers == ["WKSA ID", "First Name"]
    assert dataset.rows == []


def test_blank_rows_are_skipped_and_row_numbers_stay_the_sheets_own() -> None:
    second = ["000124", "Blair", "Park", "2012-05-01", "White", "ACTIVE"]
    content = student_workbook([CLEAN_ROW, ["", "", "", "", "", ""], second])

    dataset = load_tabular_dataset(content=content, filename="roster.xlsx")

    assert [number for number, _ in dataset.rows] == [2, 4]


# --- archives ----------------------------------------------------------------


def zip_bytes(members: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.mark.parametrize(
    "content",
    [b"PK\x03\x04 truncated", zip_bytes({"readme.txt": b"not a workbook"})],
    ids=["broken archive", "archive that is not a workbook"],
)
def test_a_malformed_archive_is_refused_cleanly(content: bytes) -> None:
    refuse_upload(
        ImportBatch.Kind.STUDENT_ROSTER,
        "roster.xlsx",
        content,
        "could not be read as a valid Excel workbook",
    )


def test_an_archive_that_expands_past_the_budget_is_refused() -> None:
    # Two megabytes of zeros compress to a few kilobytes.
    inflated = zip_bytes({"xl/padding.bin": bytes(2 * 1024 * 1024)})
    assert len(inflated) < 10_000

    with patch("data_imports.parsers.MAX_XLSX_EXPANDED_BYTES", 1024 * 1024):
        refuse_upload(
            ImportBatch.Kind.STUDENT_ROSTER,
            "roster.xlsx",
            inflated,
            "expands to more than 1 MiB",
        )


def test_an_archive_with_too_many_parts_is_refused() -> None:
    with patch("data_imports.parsers.MAX_XLSX_ARCHIVE_MEMBERS", 3):
        with pytest.raises(UnsupportedImportFileError, match="more than 3 parts"):
            load_tabular_dataset(
                content=student_workbook([CLEAN_ROW]),
                filename="roster.xlsx",
            )


# --- legacy workbooks --------------------------------------------------------


def cycle_sheets(count: int, rows_each: int = 1) -> bytes:
    def build(workbook: Workbook) -> None:
        for index in range(count):
            sheet = workbook.create_sheet(f"{index % 12 + 1:02d}-{2000 + index}")
            sheet.append(MAIN_TABLE_HEADER)
            for row in range(rows_each):
                sheet.append([f"Student {row}", "White", "Yellow", "Pass", ""])

    return workbook_bytes(build)


def test_a_workbook_with_too_many_sheets_is_refused() -> None:
    with (
        patch("data_imports.ksw_adapter.MAX_SHEETS", 2),
        pytest.raises(WorkbookStructureError, match="more than 2 sheets"),
    ):
        parse_ksw_colchester_workbook(cycle_sheets(3))


def test_the_row_budget_is_shared_across_every_sheet() -> None:
    # Two sheets of three rows (header included) fit six; a seventh does not.
    content = cycle_sheets(2, rows_each=2)
    with patch("data_imports.ksw_adapter.MAX_WORKBOOK_ROWS", 6):
        assert len(parse_ksw_colchester_workbook(content).rows) == 4
    with (
        patch("data_imports.ksw_adapter.MAX_WORKBOOK_ROWS", 5),
        pytest.raises(WorkbookStructureError, match="more than 5 rows"),
    ):
        parse_ksw_colchester_workbook(content)


def test_scanning_an_unknown_sheet_counts_against_the_budget() -> None:
    def build(workbook: Workbook) -> None:
        workbook.create_sheet("Scratch")["A100"] = "note"
        sheet = workbook.create_sheet("06-2025")
        sheet.append(MAIN_TABLE_HEADER)
        sheet.append(["A. Example", "White", "Yellow", "Pass", ""])

    with (
        patch("data_imports.ksw_adapter.MAX_WORKBOOK_ROWS", 50),
        pytest.raises(WorkbookStructureError, match="more than 50 rows"),
    ):
        parse_ksw_colchester_workbook(workbook_bytes(build))


@pytest.mark.parametrize(
    ("content", "budget"),
    [
        (standard_workbook([["A. Example", "White", "Yellow", "Pass", ""]]), 100),
        (cycle_sheets(2, rows_each=5), 3),
    ],
    ids=["parsed", "refused mid-read"],
)
def test_the_workbook_is_closed_whether_or_not_it_parses(
    content: bytes,
    budget: int,
) -> None:
    close: Callable[[Workbook], None] = Workbook.close
    with (
        patch("data_imports.ksw_adapter.MAX_WORKBOOK_ROWS", budget),
        patch.object(Workbook, "close", autospec=True, side_effect=close) as spy,
    ):
        try:
            parse_ksw_colchester_workbook(content)
        except WorkbookStructureError:
            pass
    assert spy.call_count == 1


def test_a_legacy_upload_is_checked_before_it_is_stored() -> None:
    refuse_upload(
        ImportBatch.Kind.LEGACY_GRADINGS,
        "workbook.xlsx",
        b"not a workbook",
        "could not be read as a valid Excel workbook",
    )


def test_a_legacy_workbook_with_nothing_to_import_is_refused_with_its_warnings() -> (
    None
):
    def build(workbook: Workbook) -> None:
        workbook.create_sheet("Information").append(["Cover sheet"])
        workbook.create_sheet("Notes").append(["Not a grading cycle"])

    before = stored_files()
    with pytest.raises(ValidationError) as refused:
        create_import_batch(
            kind=ImportBatch.Kind.LEGACY_GRADINGS,
            upload=SimpleUploadedFile("workbook.xlsx", workbook_bytes(build)),
            uploaded_by=admin_user(),
        )

    messages = refused.value.message_dict["file"]
    assert messages[0].startswith("No historical records were found")
    assert messages[1:] == ["Unknown sheet 'Notes' was not imported."]
    assert not ImportBatch.objects.exists()
    assert stored_files() == before
