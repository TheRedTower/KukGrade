"""The KSW Colchester workbook adapter - synthetic fixtures only.

Never the real historical workbook - every workbook
here is built in-memory from made-up names and values.
"""

from __future__ import annotations

import pytest
from openpyxl.workbook import Workbook

from data_imports.ksw_adapter import (
    WorkbookStructureError,
    parse_ksw_colchester_workbook,
)
from data_imports.testing import MAIN_TABLE_HEADER, standard_workbook, workbook_bytes
from gradings.models import LegacyAttendance, LegacyOutcome, LegacyRecordKind


def test_information_sheet_is_ignored() -> None:
    content = standard_workbook([["A. Example", "White", "Yellow Stripe", "Pass", ""]])
    result = parse_ksw_colchester_workbook(content)

    assert all(row.source_sheet != "Information" for row in result.rows)
    assert not any("Information" in warning for warning in result.warnings)


def test_valid_cycle_sheet_derives_year_and_month() -> None:
    content = standard_workbook(
        [["A. Example", "White", "Yellow Stripe", "Pass", ""]],
        sheet="03-2026",
    )
    result = parse_ksw_colchester_workbook(content)

    assert len(result.rows) == 1
    assert result.rows[0].cycle_year == 2026
    assert result.rows[0].cycle_month == 3
    assert result.rows[0].source_sheet == "03-2026"


def test_standard_main_table_parses_all_five_columns() -> None:
    content = standard_workbook(
        [["A. Example", "Yellow Belt", "Blue Stripe", "Pass", "Good form"]],
    )
    result = parse_ksw_colchester_workbook(content)

    row = result.rows[0]
    assert row.source_name == "A. Example"
    assert row.source_current_rank_text == "Yellow Belt"
    assert row.source_testing_for_rank_text == "Blue Stripe"
    assert row.source_outcome_text == "Pass"
    assert row.outcome == LegacyOutcome.PASS
    assert row.comments == "Good form"
    assert row.record_kind == LegacyRecordKind.GRADING_RESULT
    assert row.attendance_status == LegacyAttendance.ATTENDED
    assert row.source_locator.startswith("06-2025!R")


def test_blank_rows_within_the_main_table_do_not_stop_parsing() -> None:
    content = standard_workbook(
        [
            ["A. Example", "White", "Yellow Stripe", "Pass", ""],
            [None, None, None, None, None],
            ["B. Second", "White", "Yellow Stripe", "Fail", ""],
        ],
    )
    result = parse_ksw_colchester_workbook(content)

    names = [row.source_name for row in result.rows]
    assert names == ["A. Example", "B. Second"]


def test_auxiliary_section_after_the_main_table_is_captured() -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(MAIN_TABLE_HEADER)
        ws.append(["A. Example", "White", "Yellow Stripe", "Pass", ""])
        ws.append([None])
        ws.append(["Missing"])
        ws.append(["C. Third", "Left message with parent"])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    grading_kind = LegacyRecordKind.GRADING_RESULT
    aux_kind = LegacyRecordKind.AUXILIARY_NOTE
    main_rows = [r for r in result.rows if r.record_kind == grading_kind]
    aux_rows = [r for r in result.rows if r.record_kind == aux_kind]
    assert len(main_rows) == 1
    assert len(aux_rows) == 1
    assert aux_rows[0].source_name == "C. Third"
    assert aux_rows[0].source_section == "Missing"
    assert aux_rows[0].admin_note == "Left message with parent"


@pytest.mark.parametrize("section_name", ["Unaccounted", "Missing", "Current belt"])
def test_each_known_auxiliary_section_is_recognised(section_name: str) -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(MAIN_TABLE_HEADER)
        ws.append([section_name])
        ws.append(["D. Fourth", "some note"])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    assert len(result.rows) == 1
    assert result.rows[0].record_kind == LegacyRecordKind.AUXILIARY_NOTE
    assert result.rows[0].source_section == section_name


def test_did_not_attend_section_produces_did_not_attend_record_kind() -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(MAIN_TABLE_HEADER)
        ws.append(["Did not attend"])
        ws.append(["E. Fifth"])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    assert len(result.rows) == 1
    row = result.rows[0]
    assert row.record_kind == LegacyRecordKind.DID_NOT_ATTEND
    assert row.attendance_status == LegacyAttendance.DID_NOT_ATTEND
    assert row.outcome is None


def test_dna_in_the_main_outcome_column_is_did_not_attend_not_fail() -> None:
    content = standard_workbook([["F. Sixth", "White", "Yellow Stripe", "DNA", ""]])
    result = parse_ksw_colchester_workbook(content)

    row = result.rows[0]
    assert row.record_kind == LegacyRecordKind.DID_NOT_ATTEND
    assert row.attendance_status == LegacyAttendance.DID_NOT_ATTEND
    assert row.outcome is None
    assert row.outcome != LegacyOutcome.FAIL


def test_unknown_non_empty_sheet_produces_a_warning() -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("Notes")
        ws.append(["Something unrelated"])
        main = workbook.create_sheet("06-2025")
        main.append(MAIN_TABLE_HEADER)
        main.append(["A. Example", "White", "Yellow Stripe", "Pass", ""])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    assert any("Notes" in warning for warning in result.warnings)
    assert all(row.source_sheet != "Notes" for row in result.rows)


def test_unknown_empty_sheet_produces_no_warning() -> None:
    def build(workbook: Workbook) -> None:
        workbook.create_sheet("Blank Sheet")
        main = workbook.create_sheet("06-2025")
        main.append(MAIN_TABLE_HEADER)
        main.append(["A. Example", "White", "Yellow Stripe", "Pass", ""])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    assert not any("Blank Sheet" in warning for warning in result.warnings)


def test_duplicate_source_locator_is_rejected() -> None:
    # Two cycle sheets sharing the same title would produce identical
    # locators for their first data row - openpyxl silently renames a
    # duplicate sheet title, so simulate the collision directly instead.
    from data_imports.ksw_adapter import CanonicalLegacyRow, _check_locators_unique

    rows = [
        CanonicalLegacyRow(
            source_sheet="06-2025",
            source_section="",
            source_locator="06-2025!R2",
            cycle_year=2025,
            cycle_month=6,
            record_kind=LegacyRecordKind.GRADING_RESULT,
            source_name="A",
        ),
        CanonicalLegacyRow(
            source_sheet="06-2025",
            source_section="",
            source_locator="06-2025!R2",
            cycle_year=2025,
            cycle_month=6,
            record_kind=LegacyRecordKind.GRADING_RESULT,
            source_name="B",
        ),
    ]
    with pytest.raises(WorkbookStructureError, match="lists the same row twice"):
        _check_locators_unique(rows)


def test_comments_with_line_breaks_are_preserved() -> None:
    content = standard_workbook(
        [["A. Example", "White", "Yellow Stripe", "Pass", "Line one\nLine two"]],
    )
    result = parse_ksw_colchester_workbook(content)

    assert result.rows[0].comments == "Line one\nLine two"


def test_empty_comments_are_an_empty_string_not_none() -> None:
    row: list[object] = ["A. Example", "White", "Yellow Stripe", "Pass", None]
    content = standard_workbook([row])
    result = parse_ksw_colchester_workbook(content)

    assert result.rows[0].comments == ""


def test_free_form_outcome_is_preserved_unresolved() -> None:
    content = standard_workbook(
        [["A. Example", "White", "Yellow Stripe", "Testing? JKN Adi: No", ""]],
    )
    result = parse_ksw_colchester_workbook(content)

    row = result.rows[0]
    assert row.source_outcome_text == "Testing? JKN Adi: No"
    assert row.outcome is None
    assert row.record_kind == LegacyRecordKind.GRADING_RESULT


def test_na_outcome_is_preserved_unresolved_not_an_outcome() -> None:
    content = standard_workbook([["A. Example", "White", "Yellow Stripe", "N/A", ""]])
    result = parse_ksw_colchester_workbook(content)

    row = result.rows[0]
    assert row.source_outcome_text == "N/A"
    assert row.outcome is None


def test_uncertain_rank_with_question_mark_is_preserved_verbatim_unresolved() -> None:
    content = standard_workbook(
        [["A. Example", "White", "Blue Stripe?", "Pass", ""]],
    )
    result = parse_ksw_colchester_workbook(content)

    row = result.rows[0]
    assert row.source_testing_for_rank_text == "Blue Stripe?"
    # The adapter never resolves ranks - that is an explicit admin review
    # step (rank review workflow), not a parser step.


def test_typo_like_rank_text_is_preserved_verbatim() -> None:
    content = standard_workbook([["A. Example", "White", "DNB", "Pass", ""]])
    result = parse_ksw_colchester_workbook(content)

    assert result.rows[0].source_testing_for_rank_text == "DNB"


def test_abbreviated_and_duplicate_names_are_preserved_as_separate_rows() -> None:
    content = standard_workbook(
        [
            ["Alex P", "White", "Yellow Stripe", "Pass", ""],
            ["Alex P", "Yellow", "Yellow Stripe", "Pass", ""],
        ],
    )
    result = parse_ksw_colchester_workbook(content)

    assert len(result.rows) == 2
    assert result.rows[0].source_name == result.rows[1].source_name == "Alex P"
    # Distinct locators even for identical names/ranks.
    assert result.rows[0].source_locator != result.rows[1].source_locator


def test_sheet_with_no_recognisable_header_produces_a_warning_and_no_rows() -> None:
    def build(workbook: Workbook) -> None:
        ws = workbook.create_sheet("06-2025")
        ws.append(["Just some unrelated text"])

    content = workbook_bytes(build)
    result = parse_ksw_colchester_workbook(content)

    assert result.rows == []
    assert any(
        "no recognisable main table header" in warning for warning in result.warnings
    )


def test_malformed_file_raises_workbook_structure_error() -> None:
    with pytest.raises(WorkbookStructureError):
        parse_ksw_colchester_workbook(b"not a real workbook")
